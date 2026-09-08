from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets" / "dispatch-report"
OUTPUT = ROOT / "output"
OUTPUT.mkdir(exist_ok=True)
SOURCE = ASSETS / "alert-event-dispatch-center.png"

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(0, 125, 229)
MUTED = RGBColor(82, 103, 133)
LIGHT_BLUE = "EAF4FF"


def crop(name, box):
    image = Image.open(SOURCE)
    target = ASSETS / name
    image.crop(box).save(target)
    return target


def set_font(run, size, color=NAVY, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_header(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("ALERT & EVENT DISPATCH CENTER")
    set_font(run, 8, MUTED, bold=True)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Visual Proposal")
    set_font(run, 8, MUTED)


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, 25, NAVY, bold=True)
    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(18)
        r = p.add_run(subtitle)
        set_font(r, 12, MUTED)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, 17, BLUE, bold=True)


def add_copy(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_font(r, 10.5, NAVY)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, image_path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    set_font(r, 9, MUTED, italic=True)


def new_page(doc, heading, image_path, width, caption, copy):
    doc.add_page_break()
    add_heading(doc, heading)
    add_figure(doc, image_path, width, caption)
    add_copy(doc, copy)


def main():
    # Coordinates based on the 1680 x 943 UI mockup.
    queue = crop("01-incident-queue.png", (14, 136, 470, 930))
    details = crop("02-event-details-and-evidence.png", (482, 137, 1274, 723))
    health = crop("03-trailer-health.png", (482, 722, 1275, 855))
    history = crop("04-trailer-alert-history.png", (1284, 137, 1660, 677))
    actions = crop("05-actions.png", (482, 850, 1275, 931))

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)
    add_header(sec)
    add_footer(sec)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    add_title(doc, "Alert & Event Dispatch Center", "Visual proposal for mobile surveillance trailers")
    add_figure(doc, SOURCE, 6.5, "Unified dashboard concept")
    add_copy(doc, "A unified workspace that helps teams identify critical events, validate them with live evidence, and take the right action without losing trailer context.")

    new_page(doc, "1. Prioritized Incident Queue", queue, 2.85, "Alerts are grouped by status and priority.", "Events are organized into Unreviewed and In Progress queues. Severity labels make critical issues - such as intrusion, camera outages, and low fuel - easy to identify and address first.")
    new_page(doc, "2. Event Details & Evidence", details, 6.5, "The selected alert brings operational context and proof together.", "Each event combines customer, site, trailer, event time, live camera evidence, and map location. This gives the operator the context needed to validate an alert before taking action.")
    new_page(doc, "3. Trailer Health at a Glance", health, 6.5, "Core equipment health is visible alongside the event.", "Battery, fuel, network, storage, and temperature readings help determine whether an alert is isolated or part of a broader trailer issue.")
    new_page(doc, "4. Trailer Alert History", history, 3.25, "Recent events reveal patterns that may require service or closer attention.", "The history belongs to the trailer, not the operator. It provides an immediate view of recurring security and equipment alerts, making patterns easier to recognize.")
    new_page(doc, "5. Clear Response Actions", actions, 6.5, "A consistent set of next steps is available for every event.", "Operators can acknowledge an event, review the live view, send the alert to a mobile device, email configured recipients, or resolve the alert directly from the same workspace.")

    doc.core_properties.title = "Alert & Event Dispatch Center - Visual Proposal"
    doc.core_properties.subject = "Visual concept report"
    doc.core_properties.author = ""
    doc.save(OUTPUT / "Alert-and-Event-Dispatch-Center-Visual-Proposal.docx")


if __name__ == "__main__":
    main()
